#!/usr/bin/env python3
"""Test fixture for dashboard/test/source.test.js: what the Google Sheets connector answers for a sheet, shaped as
seen in mini-test round 2 (/root/fleet-tools/state/life-crm-lite/mini-test-results/round2.md).

    python3 dashboard/test/sheets_answer.py crm.xlsx           # {"info": get_spreadsheet payload,
                                                               #  "values": {tab title: get_values payload}}
    python3 dashboard/test/sheets_answer.py crm.xlsx --book    # [{title, rows}] with typed cells, dates as ISO text

get_values payloads: {"range", "values"}; rows as shown text, each row without its trailing empty cells, no
trailing empty rows, and no "values" key at all for an empty tab. Encoded from the .xlsx with openpyxl,
independently of src/source.js.
"""
import json
import sys
from datetime import date, datetime

from openpyxl import load_workbook


def shown(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, datetime) and (v.hour, v.minute, v.second) == (0, 0, 0):
        return v.date().isoformat()
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    return str(v)


def values_payload(ws):
    rows = [[shown(v) for v in row] for row in ws.iter_rows(values_only=True)]
    rows = [r[:max([i + 1 for i, v in enumerate(r) if v] or [0])] for r in rows]
    while rows and not rows[-1]:
        rows.pop()
    title = "'" + ws.title.replace("'", "''") + "'"
    out = {"range": f"{title}!A1:AZ2000", "majorDimension": "ROWS"}
    if rows:
        out["values"] = rows
    return out


def main(path, book=False):
    wb = load_workbook(path)
    if book:
        return [{"title": ws.title, "rows": [[shown(v) if isinstance(v, (date, datetime)) else v for v in row]
                                             for row in ws.iter_rows(values_only=True)]} for ws in wb.worksheets]
    return {"info": {"spreadsheetId": "x", "properties": {"title": "crm"}, "revisionId": "r1",
                     "sheets": [{"properties": {"sheetId": i, "title": ws.title}} for i, ws in enumerate(wb.worksheets)]},
            "values": {ws.title: values_payload(ws) for ws in wb.worksheets}}


if __name__ == "__main__":
    json.dump(main(sys.argv[1], "--book" in sys.argv[2:]), sys.stdout)
