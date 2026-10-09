#!/usr/bin/env python3
"""Convert a life-crm lite data block (JSON) to a life-crm sheet (.xlsx) and back.

    python3 tools/convert.py to-xlsx examples/maya/data.json examples/maya/crm.xlsx
    python3 tools/convert.py to-json examples/maya/crm.xlsx [out.json]   (stdout without out.json)
    python3 tools/convert.py check [data.json ...]   validate + round trip (default: examples/*/data.json)
    python3 tools/convert.py blank templates/blank.xlsx

The sheet layout is documented in docs/sheet-format.md and the data block in
docs/data-contract.md. The sheet opens in the full life-crm kit unchanged (apps/build.py).
Requirements: Python 3.10+, openpyxl.
"""
from __future__ import annotations

import json
import re
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent

TIMELINE_COLS = ["id", "name", "group", "color", "goal", "description", "link", "show", "order"]
STEP_COLS = ["timeline", "track", "title", "kind", "start", "end", "date", "status", "progress", "owner",
             "phase", "pin", "notes", "link", "show", "id", "importance", "urgency", "repeat"]
SETTINGS_COLS = ["key", "value", "meaning"]
COLLECTION_COLS = ["id", "name", "tab", "layout", "title_field", "status_field", "statuses", "date_field",
                   "fields", "group", "icon", "description", "empty_text", "show", "order", "area"]
PEOPLE_COLS = ["name", "role", "area", "contact"]

# Data-block keys stored in the sheet's own columns; any other key travels as an extra column.
AREA_KEYS = {"id", "name", "color", "goal", "why"}
STEP_KEYS = {"id", "area", "title", "status", "owner", "date", "start", "end", "repeat", "importance",
             "urgency", "notes", "link"}
LIST_KEYS = {"id", "name", "area", "columns", "rows"}
PERSON_KEYS = {"name", "role", "area", "contact"}
# Sheet columns an extra field may fill when it has exactly that column's name: the ones the converter
# leaves empty and the derived ones it never reads back (a full-kit `kind: milestone`, a list's `group`).
# Every other column holds a data-block value (title, role…) or steers reading (a list's tab, title_field,
# fields), so an extra field there would overwrite it.
OPEN_COLS = {"Timelines": {"group", "link", "show", "order"},
             "Steps": {"track", "kind", "progress", "phase", "pin", "show"},
             "Collections": {"layout", "status_field", "statuses", "date_field", "group", "icon", "description",
                             "empty_text", "show", "order"},
             "People": set()}

# Settings rows: (sheet key, data-block key, meaning shown in the sheet).
SETTINGS = [
    ("lite", "lite", "Version of the life-crm lite format. Leave it at 1."),
    ("title", "title", "Title shown at the top of the page"),
    ("name", "owner", "Your name"),
    ("updated", "updated", "Date of the last change (YYYY-MM-DD)"),
    ("tone", "tone", "How the assistant talks to you: kick, caring, motivational or none"),
    ("tone_line", "toneLine", "One line in that tone, shown at the top of the page (empty for none)"),
    ("sheet_url", "sheet.url", "Link to this sheet"),
    ("account", "sheet.account", "Google account this sheet is saved in"),
]
SETTING_ALIASES = {"toneline": "tone_line"}
SHEET_KEYS = {"url", "account"}
# Extra fields in Settings: plain text stays under its own key (like the full kit's window_start);
# anything else is JSON text under "json:<key>", and extra `sheet` fields under "json:sheet.<key>".
JSON_TAG = "json:"
SHEET_TAG = "sheet."
TOP_KEYS = {"lite", "title", "owner", "updated", "tone", "toneLine", "sheet", "areas", "steps", "lists", "people"}

RESERVED_TABS = {"timelines", "steps", "settings", "collections", "people", "profile", "entries"}
PEOPLE_ID = "people"
STATUSES = {"todo", "doing", "waiting", "stuck", "done"}
# A list's `layout` (its Collections `layout` cell): how its page looks. Not set = table. The full kit
# (life-crm apps/lifecrm/dashboard.py) knows table, cards and feed; it shows a board as a table, with a warning.
LAYOUTS = ["table", "cards", "board", "feed"]
LEVELS = {"high", "low"}
TONES = {"kick", "caring", "motivational", "none"}
COLORS = {"teal", "indigo", "amber", "blue", "pink", "green", "violet", "red", "orange", "gray"}
REPEAT_RE = re.compile(r"daily|weekly|monthly|every \d+ (days?|weeks?|months?)")
DATE_RE = re.compile(r"\d{4}-\d{2}(-\d{2})?")
DATE_COLUMN_RE = re.compile(r"date|deadline|due|renew")


# ---------- small helpers ----------

def header_key(cell: object) -> str:
    """Same rule as life-crm apps/lifecrm/sheet.py: 'Due date (YYYY-MM-DD)' -> 'due_date'."""
    return re.sub(r"\s+", "_", text(cell).split("(")[0].strip().lower())


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def cell_value(value: object) -> object:
    """A cell as a data-block value: text and numbers as they are, sheet dates as ISO text."""
    if isinstance(value, (datetime, date)):
        return text(value)
    return value


def or_none(value: object) -> str | None:
    s = text(value).strip()
    return s or None


def to_cell(value: object) -> object:
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def tab_named(wb, name: str):
    """Tab by name, case-insensitive, like life-crm; None if missing."""
    for ws in wb.worksheets:
        if ws.title.strip().lower() == name.strip().lower():
            return ws
    return None


def date_column(columns: list[str]) -> str:
    return next((c for c in columns if DATE_COLUMN_RE.search(header_key(c))), "")


def real_date(value: str) -> bool:
    """A real calendar day (YYYY-MM-DD) or month (YYYY-MM)."""
    if not DATE_RE.fullmatch(value):
        return False
    try:
        date.fromisoformat(value if len(value) == 10 else value + "-01")
    except ValueError:
        return False
    return True


# ---------- data block -> rows ----------

def area_row(a: dict, i: int) -> dict:
    return {"id": a.get("id"), "name": a.get("name"), "group": None, "color": a.get("color"),
            "goal": a.get("goal"), "description": a.get("why"), "show": "yes", "order": i + 1}


def step_row(s: dict) -> dict:
    start, end = s.get("start"), s.get("end")
    return {"timeline": s.get("area"), "title": s.get("title"), "kind": "period" if start and end else "task",
            "start": start, "end": end, "date": s.get("date"), "status": s.get("status"),
            "owner": s.get("owner"), "notes": s.get("notes"), "link": s.get("link"), "show": "yes",
            "id": s.get("id"), "importance": s.get("importance"), "urgency": s.get("urgency"),
            "repeat": s.get("repeat")}


def list_row(lst: dict, i: int, tab: str, area_names: dict) -> dict:
    columns = [str(c) for c in lst.get("columns") or []]
    status = next((c for c in columns if header_key(c) == "status"), "")
    return {"id": lst.get("id"), "name": lst.get("name"), "tab": tab, "layout": "table",
            "title_field": columns[0] if columns else None, "status_field": status or None,
            "date_field": date_column(columns) or None, "fields": ", ".join(columns[1:]) or None,
            "group": area_names.get(lst.get("area")) or "Lists", "show": "yes", "order": i + 1,
            "area": lst.get("area")}


def people_collection_row(order: int, tab: str) -> dict:
    return {"id": PEOPLE_ID, "name": "People", "tab": tab, "layout": "table", "title_field": "name",
            "fields": "role, area, contact", "group": "People", "show": "yes", "order": order}


def person_row(p: dict) -> dict:
    return {k: p.get(k) for k in PEOPLE_COLS}


def list_tab_names(lists: list[dict]) -> list[str]:
    """One tab per list, named after the list (at most 31 characters, no []:*?/\\, unique)."""
    used, out = set(RESERVED_TABS), []
    for lst in lists:
        base = re.sub(r"[\[\]:*?/\\]", "", str(lst.get("name") or lst.get("id") or "List")).strip()[:31] or "List"
        name, n = base, 1
        if name.lower() in used:
            name = f"{base[:26]} list"
        while name.lower() in used:
            n += 1
            name = f"{base[:26]} ({n})"
        used.add(name.lower())
        out.append(name)
    return out


# ---------- workbook writing ----------

HEAD_FILL = PatternFill("solid", fgColor="E8F0EE")


def write_tab(wb, title: str, cols: list[str], rows: list[tuple[dict, dict]]) -> None:
    """rows: (row by column, extra keys). An extra key fills the column of the same name, or gets its own."""
    cols = list(cols)
    for _, extras in rows:
        for k in extras:
            if header_key(k) not in {header_key(c) for c in cols}:
                cols.append(k)
    index = {header_key(c): n for n, c in enumerate(cols)}
    ws = wb.create_sheet(title)
    ws.append(cols)
    keep_text(ws[1])
    for row, extras in rows:
        values = [None] * len(cols)
        for k, v in row.items():
            values[index[header_key(k)]] = v
        for k, v in extras.items():
            values[index[header_key(k)]] = v
        ws.append([None if v is None or v == "" else to_cell(v) for v in values])
        keep_text(ws[ws.max_row])
    style(ws)


def keep_text(cells) -> None:
    """A value starting with '=' stays text, never a formula (headers included)."""
    for c in cells:
        if isinstance(c.value, str) and c.value.startswith("="):
            c.data_type = "s"


def style(ws) -> None:
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), HEAD_FILL
    ws.freeze_panes = "A2"
    for col in ws.columns:
        longest = max(len(text(c.value)) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, longest + 2), 60)
        for c in col[1:]:
            c.alignment = Alignment(vertical="top", wrap_text=longest > 60)


def extras_of(obj: dict, known: set[str]) -> dict:
    return {k: v for k, v in obj.items() if k not in known}


def to_workbook(data: dict) -> Workbook:
    areas, steps = data.get("areas") or [], data.get("steps") or []
    lists, people = data.get("lists") or [], data.get("people") or []
    area_names = {a.get("id"): a.get("name") for a in areas}
    tabs = list_tab_names(lists)

    wb = Workbook()
    wb.remove(wb.active)
    write_tab(wb, "Timelines", TIMELINE_COLS, [(area_row(a, i), extras_of(a, AREA_KEYS)) for i, a in enumerate(areas)])
    write_tab(wb, "Steps", STEP_COLS, [(step_row(s), extras_of(s, STEP_KEYS)) for s in steps])

    sheet = data.get("sheet") or {}
    settings = []
    for key, path, meaning in SETTINGS:
        top, _, sub = path.partition(".")
        value = sheet.get(sub) if sub else data.get(top, 1 if top == "lite" else None)
        settings.append(({"key": key, "value": value, "meaning": meaning}, {}))
    settings += extra_settings_rows(data)
    write_tab(wb, "Settings", SETTINGS_COLS, settings)

    collections = [(list_row(lst, i, tabs[i], area_names), extras_of(lst, LIST_KEYS)) for i, lst in enumerate(lists)]
    collections.append((people_collection_row(len(lists) + 1, "People"), {}))
    write_tab(wb, "Collections", COLLECTION_COLS, collections)
    write_tab(wb, "People", PEOPLE_COLS, [(person_row(p), extras_of(p, PERSON_KEYS)) for p in people])

    for lst, tab in zip(lists, tabs):
        columns = [str(c) for c in lst.get("columns") or []]
        rows = []
        for r in lst.get("rows") or []:
            row = {c: r.get(c) for c in columns}
            rows.append((row, {k: v for k, v in r.items() if k not in columns}))
        write_tab(wb, tab, columns, rows)
    return wb


def setting_key(key: str) -> str:
    k = key.strip().lower()
    return SETTING_ALIASES.get(k, k)


def extra_settings_rows(data: dict) -> list[tuple[dict, dict]]:
    """Unknown top-level fields and extra `sheet` fields, as Settings rows (see JSON_TAG)."""
    known = {key for key, _, _ in SETTINGS}
    rows = []
    for k, v in data.items():
        if k in TOP_KEYS:
            continue
        plain = (isinstance(v, str) and k == k.strip() and setting_key(k) not in known
                 and not k.lower().startswith(JSON_TAG))
        rows.append(({"key": k, "value": v} if plain
                     else {"key": JSON_TAG + k, "value": json.dumps(v, ensure_ascii=False)}, {}))
    sheet = data.get("sheet")
    for k, v in sheet.items() if isinstance(sheet, dict) else []:
        if k not in SHEET_KEYS:
            rows.append(({"key": JSON_TAG + SHEET_TAG + k, "value": json.dumps(v, ensure_ascii=False)}, {}))
    return rows


def json_or_text(value: object) -> object:
    try:
        return json.loads(text(value))
    except ValueError:
        return text(value)


# ---------- workbook reading ----------

def read_tab(ws) -> tuple[list[str], list[dict]]:
    """(raw headers, rows as {raw header: cell value}) for every non-empty row under the header."""
    it = ws.iter_rows(values_only=True)
    headers = [text(c).strip() for c in next(it, [])]
    rows = []
    for row in it:
        item = {h: cell_value(v) for h, v in zip(headers, row) if h}
        if any(text(v).strip() for v in item.values()):
            rows.append(item)
    return [h for h in headers if h], rows


def take_extras(obj: dict, raw: dict, mapped: set[str], derived: dict, base_cols: list[str]) -> dict:
    """Keep every non-empty cell that the data block doesn't map and the converter didn't derive."""
    for h, v in raw.items():
        k = header_key(h)
        if k in mapped or not text(v).strip():
            continue
        if k in derived and text(v).strip() == text(derived[k]):
            continue
        obj[k if k in base_cols else h] = v
    return obj


def by_key(raw: dict) -> dict:
    return {header_key(h): v for h, v in raw.items()}


def from_workbook(wb) -> dict:
    # Settings (and Profile's name, for sheets made with the full kit)
    settings, extra_settings, sheet_extras = {}, {}, {}
    ws = tab_named(wb, "Settings")
    known = {key for key, _, _ in SETTINGS}
    for row in ws.iter_rows(min_row=2, values_only=True) if ws else []:
        if len(row) >= 2 and text(row[0]).strip():
            raw_key = text(row[0]).strip()
            key = setting_key(raw_key)
            if key.startswith(JSON_TAG):
                name = raw_key[len(JSON_TAG):]
                if name.lower().startswith(SHEET_TAG):
                    sheet_extras[name[len(SHEET_TAG):]] = json_or_text(row[1])
                else:
                    extra_settings[name] = json_or_text(row[1])
            elif key in known:
                settings[key] = cell_value(row[1])
            else:
                extra_settings[raw_key] = text(row[1])
    profile = tab_named(wb, "Profile")
    if not text(settings.get("name")) and profile:
        for row in profile.iter_rows(min_row=2, values_only=True):
            if len(row) >= 2 and text(row[0]).strip().lower() == "name":
                settings["name"] = row[1]
    lite = text(settings.get("lite")).strip()
    data = {
        "lite": int(lite) if lite.isdigit() else 1,
        "title": text(settings.get("title")),
        "owner": text(settings.get("name")),
        "updated": text(settings.get("updated")),
        "tone": text(settings.get("tone")).strip().lower() or "none",
        "toneLine": text(settings.get("tone_line")),
        "sheet": {"url": text(settings.get("sheet_url")), "account": text(settings.get("account")), **sheet_extras},
    }

    # Timelines -> areas
    areas = []
    ws = tab_named(wb, "Timelines")
    for i, raw in enumerate(read_tab(ws)[1] if ws else []):
        r = by_key(raw)
        a = {"id": text(r.get("id")).strip(), "name": text(r.get("name")),
             "color": text(r.get("color")).strip().lower() or "gray", "goal": text(r.get("goal")),
             "why": text(r.get("description"))}
        areas.append(take_extras(a, raw, {"id", "name", "color", "goal", "description"}, area_row(a, i), TIMELINE_COLS))
    data["areas"] = areas

    # Steps
    steps = []
    ws = tab_named(wb, "Steps")
    mapped = {"timeline", "title", "start", "end", "date", "status", "owner", "notes", "link", "id",
              "importance", "urgency", "repeat"}
    for i, raw in enumerate(read_tab(ws)[1] if ws else []):
        r = by_key(raw)
        s = {"id": text(r.get("id")).strip() or f"s{i + 1}", "area": text(r.get("timeline")).strip(),
             "title": text(r.get("title")), "status": text(r.get("status")).strip().lower() or "todo",
             "owner": text(r.get("owner")).strip() or "me", "date": or_none(r.get("date")),
             "start": or_none(r.get("start")), "end": or_none(r.get("end")), "repeat": or_none(r.get("repeat")),
             "importance": (or_none(r.get("importance")) or "").lower() or None,
             "urgency": (or_none(r.get("urgency")) or "").lower() or None,
             "notes": text(r.get("notes")), "link": text(r.get("link"))}
        steps.append(take_extras(s, raw, mapped, step_row(s), STEP_COLS))
    data["steps"] = steps

    # Collections -> lists (+ the People tab)
    area_names = {a["id"]: a["name"] for a in areas}
    lists, people_tab = [], "People"
    ws = tab_named(wb, "Collections")
    for raw in read_tab(ws)[1] if ws else []:
        r = by_key(raw)
        cid = text(r.get("id")).strip().lower()
        tab = text(r.get("tab")).strip() or text(r.get("name")).strip()
        if cid == PEOPLE_ID:
            people_tab = tab or people_tab
            continue
        lst = {"id": cid, "name": text(r.get("name")) or cid.title(), "area": text(r.get("area")).strip()}
        data_ws = tab_named(wb, tab) if tab else None
        headers, rows = read_tab(data_ws) if data_ws else ([], [])
        # The list's own columns are title_field + fields; any other header is an extra field of its rows.
        named = {header_key(f) for f in [r.get("title_field"), *re.split(r"[|,]", text(r.get("fields")))]} - {""}
        columns = [h for h in headers if header_key(h) in named] if named else headers
        lst["columns"] = columns
        lst["rows"] = [{**{c: ("" if row.get(c) is None else row.get(c)) for c in columns},
                        **{h: v for h, v in row.items() if h not in columns and text(v).strip()}} for row in rows]
        lists.append(take_extras(lst, raw, {"id", "name", "tab", "area"},
                                 list_row(lst, len(lists), tab, area_names), COLLECTION_COLS))
    data["lists"] = lists

    people = []
    ws = tab_named(wb, people_tab)
    for raw in read_tab(ws)[1] if ws else []:
        r = by_key(raw)
        p = {k: text(r.get(k)) for k in PEOPLE_COLS}
        people.append(take_extras(p, raw, set(PEOPLE_COLS), {}, PEOPLE_COLS))
    data["people"] = people

    data.update(extra_settings)
    return data


# ---------- checks ----------

def validate(data: dict) -> list[str]:
    """Problems against docs/data-contract.md (empty list = valid)."""
    out = []
    if data.get("lite") != 1:
        out.append("lite must be 1")
    if data.get("tone") not in TONES:
        out.append(f"tone '{data.get('tone')}' is not one of {', '.join(sorted(TONES))}")
    area_ids = [a.get("id") for a in data.get("areas") or []]
    for a in data.get("areas") or []:
        if not a.get("id") or not a.get("name"):
            out.append(f"area {a}: needs id and name")
        if a.get("color") not in COLORS:
            out.append(f"area {a.get('id')}: color '{a.get('color')}' is not a life-crm colour")
    if len(set(area_ids)) != len(area_ids):
        out.append("area ids must be unique")
    step_ids = [s.get("id") for s in data.get("steps") or []]
    if len(set(step_ids)) != len(step_ids):
        out.append("step ids must be unique")
    for s in data.get("steps") or []:
        where = f"step {s.get('id')} '{s.get('title')}'"
        if s.get("area") not in area_ids:
            out.append(f"{where}: unknown area '{s.get('area')}'")
        if s.get("status") not in STATUSES:
            out.append(f"{where}: status '{s.get('status')}'")
        for k in ("importance", "urgency"):
            if s.get(k) not in LEVELS:
                out.append(f"{where}: {k} must be high or low")
        for k in ("date", "start", "end"):
            if s.get(k) is not None and not real_date(str(s.get(k))):
                out.append(f"{where}: {k} '{s.get(k)}' is not a real date (YYYY-MM-DD) or month (YYYY-MM)")
        if s.get("date") and (s.get("start") or s.get("end")):
            out.append(f"{where}: either date, or start + end, not both")
        if bool(s.get("start")) != bool(s.get("end")):
            out.append(f"{where}: start and end go together")
        if s.get("repeat") is not None and not REPEAT_RE.fullmatch(str(s.get("repeat"))):
            out.append(f"{where}: repeat '{s.get('repeat')}' is not a form the page can compute")
        if s.get("repeat") and not s.get("date"):
            out.append(f"{where}: a repeating step needs date (the next due date)")
        if not s.get("owner"):
            out.append(f"{where}: owner is 'me' or a person's name")
    list_ids = [lst.get("id") for lst in data.get("lists") or []]
    if len(set(list_ids)) != len(list_ids) or PEOPLE_ID in list_ids:
        out.append(f"list ids must be unique and not '{PEOPLE_ID}'")
    for lst in data.get("lists") or []:
        if lst.get("area") and lst.get("area") not in area_ids:
            out.append(f"list {lst.get('id')}: unknown area '{lst.get('area')}'")
        columns = [str(c) for c in lst.get("columns") or []]
        if not columns:
            out.append(f"list {lst.get('id')}: needs at least one column")
        if any(not header_key(c) or re.search(r"[,|]", c) for c in columns):
            out.append(f"list {lst.get('id')}: column names must not be empty or contain , or |")
        out += [f"list {lst.get('id')}: {c}" for c in key_collisions(columns, lst.get("rows") or [])]
        layout = lst.get("layout")
        if layout is not None and layout not in LAYOUTS:
            out.append(f"list {lst.get('id')}: layout must be one of {', '.join(LAYOUTS)}, not {layout!r}")
        if layout == "board" and not any(header_key(c) == "status" for c in columns):
            out.append(f"list {lst.get('id')}: a board needs a column named status (its values are the board's columns)")
        if lst.get("statuses") is not None and not isinstance(lst.get("statuses"), str):
            out.append(f"list {lst.get('id')}: statuses is text, the board's columns in order, e.g. \"idea, writing, done\"")
    for p in data.get("people") or []:
        if not p.get("name"):
            out.append(f"person {p}: needs a name")
        if p.get("area") and p.get("area") not in area_ids:
            out.append(f"person {p.get('name')}: unknown area '{p.get('area')}'")
    out += field_collisions("areas", [(f"area {a.get('id')}", a) for a in data.get("areas") or []],
                            AREA_KEYS, TIMELINE_COLS, OPEN_COLS["Timelines"])
    out += field_collisions("steps", [(f"step {s.get('id')} '{s.get('title')}'", s) for s in data.get("steps") or []],
                            STEP_KEYS, STEP_COLS, OPEN_COLS["Steps"])
    out += field_collisions("lists", [(f"list {lst.get('id')}", lst) for lst in data.get("lists") or []],
                            LIST_KEYS, COLLECTION_COLS, OPEN_COLS["Collections"])
    out += field_collisions("people", [(f"person {p.get('name')}", p) for p in data.get("people") or []],
                            PERSON_KEYS, PEOPLE_COLS, OPEN_COLS["People"])
    return out


def field_collisions(label: str, objects: list[tuple[str, dict]], known: set[str], cols: list[str],
                     open_cols: set[str]) -> list[str]:
    """Extra fields of areas, steps, lists or people that would land in a column already in use.

    An extra field may fill one of the tab's columns only if it has exactly that column's name and the
    column is in OPEN_COLS; otherwise one value would overwrite the other."""
    col_by_key = {header_key(c): c for c in cols}
    out, names = [], {}
    for where, obj in objects:
        for k in obj:
            if k in known:
                continue
            key = header_key(k)
            col = col_by_key.get(key)
            if not key:
                out.append(f"{where}: field {k!r} is empty as a sheet header")
            elif col is not None and not (k == col and col in open_cols):
                out.append(f"{where}: field {k!r} would share the sheet column '{col}'; "
                           + (f"write it as '{col}' or rename it" if col in open_cols
                              else "that column already holds another value, so rename it"))
            elif col is None:
                same = names.setdefault(key, [])
                if k not in same:
                    same.append(k)
    out += [f"{label}: {', '.join(map(repr, same))} would share one sheet column ('{key}'); "
            f"rename them so they differ in more than case, spaces or anything in brackets"
            for key, same in names.items() if len(same) > 1]
    return out


def key_collisions(columns: list[str], rows: list[dict]) -> list[str]:
    """Column names and row keys of one list that become the same sheet column (header_key)."""
    names: dict[str, list[str]] = {}
    for name in [*columns, *(str(k) for row in rows for k in row)]:
        same = names.setdefault(header_key(name), [])
        if name not in same:
            same.append(name)
    out = [f"row key {k!r} is empty as a sheet header" for k in names.get("", []) if k not in columns]
    out += [f"{', '.join(map(repr, same))} would share one sheet column ('{key}'); "
            f"rename them so they differ in more than case, spaces or anything in brackets"
            for key, same in names.items() if key and len(same) > 1]
    out += [f"column {c!r} is listed twice" for c in dict.fromkeys(c for c in columns if columns.count(c) > 1)]
    return out


def diff(a: object, b: object, path: str = "") -> list[str]:
    if isinstance(a, dict) and isinstance(b, dict):
        return [d for k in sorted(set(a) | set(b), key=str)
                for d in (diff(a.get(k, "<missing>"), b.get(k, "<missing>"), f"{path}.{k}"))]
    if isinstance(a, list) and isinstance(b, list) and len(a) == len(b):
        return [d for i, (x, y) in enumerate(zip(a, b)) for d in diff(x, y, f"{path}[{i}]")]
    return [] if a == b else [f"{path or '.'}: {a!r} -> {b!r}"]


def check(paths: list[Path]) -> int:
    failed = 0
    with tempfile.TemporaryDirectory() as tmp:
        for path in paths:
            data = json.loads(path.read_text())
            problems = validate(data)
            xlsx = Path(tmp) / f"{path.parent.name}.xlsx"
            to_workbook(data).save(xlsx)
            back = from_workbook(load_workbook(xlsx))
            changes = diff(data, back)
            name = path.relative_to(ROOT) if path.is_relative_to(ROOT) else path
            if problems or changes:
                failed += 1
                print(f"FAIL {name}")
                for p in problems:
                    print("  invalid:", p)
                for c in changes:
                    print("  round trip changed", c)
            else:
                print(f"ok   {name}: valid; data.json -> xlsx -> data.json identical "
                      f"({len(data['areas'])} areas, {len(data['steps'])} steps, "
                      f"{len(data['lists'])} lists, {len(data['people'])} people)")
    return 1 if failed else 0


BLANK = {"lite": 1}


def main(argv: list[str]) -> int:
    cmd, args = (argv[0], argv[1:]) if argv else ("check", [])
    if cmd == "to-xlsx" and len(args) == 2:
        data = json.loads(Path(args[0]).read_text())
        for p in validate(data):
            print("warning:", p, file=sys.stderr)
        Path(args[1]).parent.mkdir(parents=True, exist_ok=True)
        to_workbook(data).save(args[1])
        print(f"wrote {args[1]}")
        return 0
    if cmd == "to-json" and len(args) in (1, 2):
        out = json.dumps(from_workbook(load_workbook(args[0])), ensure_ascii=False, indent=2) + "\n"
        if len(args) == 2:
            Path(args[1]).write_text(out)
            print(f"wrote {args[1]}")
        else:
            sys.stdout.write(out)
        return 0
    if cmd == "check":
        paths = [Path(p).resolve() for p in args] or sorted(ROOT.glob("examples/*/data.json"))
        if not paths:
            print("no data.json files to check", file=sys.stderr)
            return 1
        return check(paths)
    if cmd == "blank" and len(args) == 1:
        Path(args[0]).parent.mkdir(parents=True, exist_ok=True)
        to_workbook(BLANK).save(args[0])
        print(f"wrote {args[0]}")
        return 0
    print(__doc__, file=sys.stderr)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
