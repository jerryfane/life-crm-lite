#!/usr/bin/env python3
"""Tests for tools/convert.py:  python3 tools/test_convert.py  (stdlib unittest + openpyxl)."""
from __future__ import annotations

import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import convert  # noqa: E402

MAYA = json.loads((convert.ROOT / "examples" / "maya" / "data.json").read_text())


def maya() -> dict:
    return copy.deepcopy(MAYA)


def save(data: dict, folder: str) -> Path:
    path = Path(folder) / "crm.xlsx"
    convert.to_workbook(data).save(path)
    return path


def round_trip(data: dict) -> dict:
    with tempfile.TemporaryDirectory() as tmp:
        return convert.from_workbook(load_workbook(save(data, tmp)))


class Dates(unittest.TestCase):
    def test_impossible_dates_and_months_are_invalid(self):
        for bad in ["2026-13", "2026-00", "2026-13-99", "2026-02-30", "2026-04-31"]:
            data = maya()
            data["steps"][0]["date"], data["steps"][0]["start"], data["steps"][0]["end"] = bad, None, None
            self.assertTrue(any("real date" in p for p in convert.validate(data)), bad)

    def test_real_dates_and_months_are_valid(self):
        for good in ["2026-11", "2028-02-29", "2026-12-31"]:
            data = maya()
            data["steps"][0]["date"], data["steps"][0]["start"], data["steps"][0]["end"] = good, None, None
            self.assertEqual(convert.validate(data), [], good)


class UnknownFields(unittest.TestCase):
    def test_unknown_top_level_and_sheet_fields_round_trip(self):
        data = maya()
        data["sheet"]["folder"] = "https://drive.google.com/drive/folders/EXAMPLE"
        data["sheet"]["tabs"] = {"steps": 2, "hidden": False}
        data["coach"] = {"style": "kick", "days": ["mon", "thu"]}
        data["streak"] = 3
        data["beta"] = True
        data["nothing"] = None
        data["note"] = "plain text"
        data["Title"] = "clashes with the title setting"
        data["json:odd"] = "starts with the tag"
        self.assertEqual(convert.diff(data, round_trip(data)), [])

    def test_plain_text_setting_stays_readable_by_the_full_kit(self):
        # Guard: life-crm reads its own settings (window_start…) as plain key/value rows.
        data = maya()
        data["window_start"] = "2026-01"
        with tempfile.TemporaryDirectory() as tmp:
            ws = load_workbook(save(data, tmp))["Settings"]
            self.assertIn(("window_start", "2026-01"), [(r[0], r[1]) for r in ws.iter_rows(values_only=True)])

    def test_extra_list_row_fields_round_trip(self):
        data = maya()
        rows = data["lists"][0]["rows"]
        rows[0]["contact"] = "program office"
        rows[1]["visa"] = "J-1"
        rows[1]["interviews"] = 2
        self.assertEqual(convert.validate(data), [])
        back = round_trip(data)
        self.assertEqual(convert.diff(data, back), [])
        self.assertEqual(back["lists"][0]["columns"], MAYA["lists"][0]["columns"])


class CollidingKeys(unittest.TestCase):
    def test_reviewer_case_row_keys_colliding_with_a_column_or_each_other_are_invalid(self):
        # Issue #16: these used to pass validation, then 'PROGRAM' overwrote 'program' and both extras vanished.
        data = maya()
        row = data["lists"][0]["rows"][0]
        row["PROGRAM"] = "INJECTED"
        row["visa type"] = "J-1"
        row["visa_type"] = "H-1B"
        problems = convert.validate(data)
        self.assertTrue(any("'program', 'PROGRAM'" in p for p in problems), problems)
        self.assertTrue(any("'visa type', 'visa_type'" in p for p in problems), problems)

    def test_same_extra_spelled_differently_on_two_rows_is_invalid(self):
        data = maya()
        data["lists"][0]["rows"][0]["Visa"] = "J-1"
        data["lists"][0]["rows"][1]["visa"] = "J-1"
        self.assertTrue(any("'Visa', 'visa'" in p for p in convert.validate(data)))

    def test_row_key_that_is_empty_as_a_header_is_invalid(self):
        data = maya()
        data["lists"][0]["rows"][0]["(note)"] = "x"
        self.assertTrue(any("'(note)' is empty" in p for p in convert.validate(data)))

    def test_list_columns_colliding_with_each_other_are_invalid(self):
        for columns in (["program", "Program"], ["visa type", "visa_type"], ["program", "program"]):
            data = maya()
            data["lists"][0]["columns"] = columns
            data["lists"][0]["rows"] = [{c: "x" for c in columns}]
            self.assertTrue(any("share one sheet column" in p or "listed twice" in p
                                for p in convert.validate(data)), columns)

    def test_distinct_extras_on_different_rows_stay_valid(self):
        data = maya()
        data["lists"][0]["rows"][0]["contact"] = "office"
        data["lists"][0]["rows"][1]["contact"] = "dean"
        data["lists"][0]["rows"][2]["visa"] = "J-1"
        self.assertEqual(convert.validate(data), [])
        self.assertEqual(convert.diff(data, round_trip(data)), [])


class CollidingFields(unittest.TestCase):
    def test_step_title_and_person_role_extras_are_invalid(self):
        # Before: valid, then 'Title' overwrote the step's title and 'Role' the person's role, and both vanished.
        data = maya()
        data["steps"][0]["Title"] = "INJECTED"
        data["people"][0]["Role"] = "X"
        problems = convert.validate(data)
        self.assertIn("step s1 'Urology rotation (Surgery II)': field 'Title' would share the sheet column 'title'; "
                      "that column already holds another value, so rename it", problems)
        self.assertIn("person Dr. Navarro: field 'Role' would share the sheet column 'role'; "
                      "that column already holds another value, so rename it", problems)

    def test_extras_on_columns_that_hold_or_steer_data_are_invalid(self):
        data = maya()
        data["areas"][0]["description"] = "x"   # the area's `why` goes there
        data["steps"][0]["timeline"] = "x"      # the step's `area` goes there
        data["lists"][0]["tab"] = "Other"       # says which tab holds the rows
        data["lists"][1]["fields"] = "type"     # says which headers are the list's columns
        problems = convert.validate(data)
        for where, field in [("area school", "description"), ("step s1", "timeline"), ("list programs", "tab"),
                             ("list papers", "fields")]:
            self.assertTrue(any(p.startswith(where) and f"field '{field}'" in p for p in problems), (where, problems))

    def test_open_column_needs_its_exact_name(self):
        data = maya()
        data["steps"][0]["Track"] = "Rotations"
        self.assertTrue(any("write it as 'track' or rename it" in p for p in convert.validate(data)))

    def test_extras_spelled_two_ways_across_objects_are_invalid(self):
        data = maya()
        data["steps"][0]["Visa"] = "J-1"
        data["steps"][1]["visa"] = "J-1"
        self.assertTrue(any(p.startswith("steps: 'Visa', 'visa' would share one sheet column")
                            for p in convert.validate(data)))

    def test_full_kit_columns_with_their_exact_names_stay_valid(self):
        data = maya()
        data["areas"][0]["link"] = "https://drive.example/folder"
        data["areas"][1]["group"] = "Career"
        data["steps"][0]["track"] = "Rotations"
        data["steps"][1]["progress"] = 40
        data["steps"][2]["kind"] = "milestone"
        data["lists"][0]["icon"] = "building"
        data["lists"][1]["group"] = "Work"
        data["people"][0]["email"] = "n@example.com"
        self.assertEqual(convert.validate(data), [])
        self.assertEqual(convert.diff(data, round_trip(data)), [])


class Formulas(unittest.TestCase):
    def test_header_and_cells_starting_with_equals_stay_text(self):
        data = maya()
        lst = data["lists"][1]
        lst["columns"].append("=1+1")
        for row in lst["rows"]:
            row["=1+1"] = "=SUM(A1:A3)"
        data["lists"][0]["rows"][0]["=HYPERLINK(\"x\")"] = "extra"
        with tempfile.TemporaryDirectory() as tmp:
            wb = load_workbook(save(data, tmp))
            for ws in wb.worksheets:
                formulas = [c.coordinate for row in ws.iter_rows() for c in row if c.data_type == "f"]
                self.assertEqual(formulas, [], ws.title)
            self.assertEqual(convert.diff(data, convert.from_workbook(wb)), [])


class Layouts(unittest.TestCase):
    def papers(self, data: dict) -> dict:
        return next(lst for lst in data["lists"] if lst["id"] == "papers")

    def test_board_with_statuses_round_trips_through_the_collections_row(self):
        data = maya()
        self.papers(data).update(layout="board", statuses="idea, writing, submitted, accepted")
        self.assertEqual(convert.validate(data), [])
        with tempfile.TemporaryDirectory() as tmp:
            wb = load_workbook(save(data, tmp))
            head = [c.value for c in wb["Collections"][1]]
            row = next(r for r in wb["Collections"].iter_rows(min_row=2, values_only=True) if r[0] == "papers")
            self.assertEqual(row[head.index("layout")], "board")
            self.assertEqual(row[head.index("statuses")], "idea, writing, submitted, accepted")
            self.assertEqual(row[head.index("status_field")], "status")
            self.assertEqual(convert.diff(data, convert.from_workbook(wb)), [])

    def test_every_layout_round_trips_and_table_is_the_default(self):
        for layout in ["cards", "board", "feed"]:
            data = maya()
            self.papers(data)["layout"] = layout
            self.assertEqual(convert.diff(data, round_trip(data)), [], layout)
        data = maya()
        self.papers(data)["layout"] = "table"
        self.assertNotIn("layout", self.papers(round_trip(data)))  # the default: written, not brought back
        self.assertEqual(convert.diff(maya(), round_trip(maya())), [])

    def test_unknown_layout_board_without_status_and_list_statuses_are_invalid(self):
        data = maya()
        self.papers(data)["layout"] = "kanban"
        self.assertTrue(any("layout must be one of table, cards, board, feed" in p for p in convert.validate(data)))
        data = maya()
        self.papers(data).update(layout="board", columns=["paper", "type", "supervisor"])
        self.assertTrue(any("a board needs a column named status" in p for p in convert.validate(data)))
        data = maya()
        self.papers(data)["statuses"] = ["idea", "writing"]
        self.assertTrue(any("statuses is text" in p for p in convert.validate(data)))



if __name__ == "__main__":
    unittest.main(verbosity=2)
